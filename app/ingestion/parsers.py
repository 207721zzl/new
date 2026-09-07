"""把常见办公文件解析成统一的 RAG 知识文档。"""

from __future__ import annotations

import hashlib
import html
import mimetypes
import re
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from app.rag.documents import load_documents, normalize_text
from app.rag.models import KnowledgeDocument, SourceBlock


SUPPORTED_EXTENSIONS = {
    ".json",
    ".pdf",
    ".docx",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".txt",
    ".xlsx",
    ".xlsm",
    ".xls",
}


def _file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        while chunk := file.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_id(value: str, suffix: str = "") -> str:
    """由逻辑文件名生成稳定身份，使同名文档更新可以替换旧索引。"""
    normalized = re.sub(r"[^a-zA-Z0-9_-]+", "_", value).strip("_").lower()
    normalized = normalized[:72] or "document"
    extra = f"_{suffix}" if suffix else ""
    identity = f"{value.casefold()}::{suffix}"
    identity_hash = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:12]
    return f"{normalized}{extra}_{identity_hash}"[:128]


def _source_date(path: Path, candidate: datetime | None = None) -> date:
    if candidate is not None:
        return candidate.date()
    return datetime.fromtimestamp(path.stat().st_mtime, tz=UTC).date()


def _mime_type(filename: str) -> str:
    return mimetypes.guess_type(filename)[0] or "application/octet-stream"


def _make_document(
    path: Path,
    *,
    original_filename: str,
    digest: str,
    title: str,
    blocks: list[SourceBlock],
    doc_type: str,
    version: str = "1",
    updated_at: date | None = None,
    id_suffix: str = "",
    metadata: dict[str, Any] | None = None,
) -> KnowledgeDocument:
    if not blocks:
        raise ValueError(f"no readable content found in {original_filename}")
    content = "\n\n".join(block.text for block in blocks)
    return KnowledgeDocument(
        document_id=_safe_id(original_filename, id_suffix),
        title=(title.strip() or Path(original_filename).stem)[:256],
        content=content,
        blocks=blocks,
        doc_type=doc_type,
        source=original_filename,
        version=version[:32] or "1",
        updated_at=updated_at or _source_date(path),
        original_filename=original_filename,
        mime_type=_mime_type(original_filename),
        content_sha256=digest,
        metadata=metadata or {},
    )


def _decode_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"unsupported text encoding: {path.name}")


def _split_plain_paragraphs(text: str) -> list[SourceBlock]:
    return [
        SourceBlock(text=item, block_type="paragraph")
        for item in re.split(r"\n\s*\n", normalize_text(text))
        if item.strip()
    ]


def _table_markdown(rows: list[list[Any]]) -> str:
    normalized = [
        [str(value if value is not None else "").strip().replace("|", "\\|") for value in row]
        for row in rows
    ]
    normalized = [row for row in normalized if any(row)]
    if not normalized:
        return ""
    width = max(len(row) for row in normalized)
    normalized = [row + [""] * (width - len(row)) for row in normalized]
    header = normalized[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return "\n".join(lines)


def _parse_json(path: Path) -> list[KnowledgeDocument]:
    documents = load_documents(path)
    digest = _file_digest(path)
    return [
        document.model_copy(
            update={
                "original_filename": document.original_filename or path.name,
                "mime_type": document.mime_type or "application/json",
                "content_sha256": document.content_sha256 or digest,
            }
        )
        for document in documents
    ]


def _parse_txt(path: Path, original_filename: str) -> list[KnowledgeDocument]:
    digest = _file_digest(path)
    return [
        _make_document(
            path,
            original_filename=original_filename,
            digest=digest,
            title=Path(original_filename).stem,
            blocks=_split_plain_paragraphs(_decode_text(path)),
            doc_type="text",
        )
    ]


def _parse_markdown(path: Path, original_filename: str) -> list[KnowledgeDocument]:
    digest = _file_digest(path)
    text = _decode_text(path)
    blocks: list[SourceBlock] = []
    headings: list[str] = []
    first_heading: str | None = None
    buffer: list[str] = []
    in_code = False

    def flush(block_type: str = "paragraph") -> None:
        nonlocal buffer
        value = normalize_text("\n".join(buffer))
        if value:
            blocks.append(
                SourceBlock(
                    text=value,
                    block_type=block_type,
                    section_path=list(headings),
                )
            )
        buffer = []

    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            if in_code:
                buffer.append(line)
                flush("code")
            else:
                flush()
                buffer.append(line)
            in_code = not in_code
            continue
        if in_code:
            buffer.append(line)
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading_match:
            flush()
            level = len(heading_match.group(1))
            heading = heading_match.group(2).strip()
            if first_heading is None:
                first_heading = heading
            headings = headings[: level - 1] + [heading]
            continue
        if not line.strip():
            flush("table" if buffer and buffer[0].lstrip().startswith("|") else "paragraph")
            continue
        buffer.append(line)
    flush("code" if in_code else "paragraph")
    title = first_heading or Path(original_filename).stem
    return [
        _make_document(
            path,
            original_filename=original_filename,
            digest=digest,
            title=title,
            blocks=blocks,
            doc_type="markdown",
        )
    ]


def _parse_html(path: Path, original_filename: str) -> list[KnowledgeDocument]:
    try:
        from bs4 import BeautifulSoup, NavigableString, Tag
    except ImportError as exc:
        raise ValueError("HTML parsing requires beautifulsoup4") from exc

    digest = _file_digest(path)
    soup = BeautifulSoup(_decode_text(path), "html.parser")
    for tag in soup(["script", "style", "noscript", "template", "svg"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else Path(original_filename).stem
    root = soup.body or soup
    headings: list[str] = []
    blocks: list[SourceBlock] = []
    handled_tables: set[int] = set()

    for element in root.descendants:
        if isinstance(element, NavigableString) or not isinstance(element, Tag):
            continue
        name = element.name.lower()
        if name in {f"h{level}" for level in range(1, 7)}:
            level = int(name[1])
            value = element.get_text(" ", strip=True)
            if value:
                headings = headings[: level - 1] + [value]
        elif name == "table" and id(element) not in handled_tables:
            handled_tables.add(id(element))
            rows = [
                [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
                for row in element.find_all("tr")
            ]
            value = _table_markdown(rows)
            if value:
                blocks.append(SourceBlock(text=value, block_type="table", section_path=list(headings)))
        elif name in {"p", "li", "pre", "blockquote"} and not element.find_parent("table"):
            value = element.get_text(" ", strip=True)
            if value:
                block_type = {"li": "list", "pre": "code", "blockquote": "quote"}.get(name, "paragraph")
                blocks.append(SourceBlock(text=html.unescape(value), block_type=block_type, section_path=list(headings)))
    return [
        _make_document(
            path,
            original_filename=original_filename,
            digest=digest,
            title=title,
            blocks=blocks,
            doc_type="html",
        )
    ]


def _parse_docx(path: Path, original_filename: str) -> list[KnowledgeDocument]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ValueError("DOCX parsing requires python-docx") from exc

    digest = _file_digest(path)
    document = Document(path)
    properties = document.core_properties
    headings: list[str] = []
    blocks: list[SourceBlock] = []
    title = properties.title.strip() if properties.title else Path(original_filename).stem

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            value = item.text.strip()
            if not value:
                continue
            style = item.style.name or ""
            match = re.search(r"(?:Heading|标题)\s*(\d+)", style, flags=re.IGNORECASE)
            if match:
                level = min(6, max(1, int(match.group(1))))
                headings = headings[: level - 1] + [value]
                if level == 1 and title == Path(original_filename).stem:
                    title = value
                continue
            blocks.append(SourceBlock(text=value, block_type="paragraph", section_path=list(headings)))
        elif isinstance(item, Table):
            value = _table_markdown([[cell.text for cell in row.cells] for row in item.rows])
            if value:
                blocks.append(SourceBlock(text=value, block_type="table", section_path=list(headings)))

    version = properties.version or (str(properties.revision) if properties.revision else "1")
    return [
        _make_document(
            path,
            original_filename=original_filename,
            digest=digest,
            title=title,
            blocks=blocks,
            doc_type="word",
            version=version,
            updated_at=_source_date(path, properties.modified),
            metadata={"author": properties.author or ""},
        )
    ]


def _parse_pdf(
    path: Path,
    original_filename: str,
    *,
    enable_ocr: bool,
    ocr_language: str,
    ocr_dpi: int,
) -> list[KnowledgeDocument]:
    try:
        import pymupdf
    except ImportError as exc:
        raise ValueError("PDF parsing requires PyMuPDF") from exc

    digest = _file_digest(path)
    blocks: list[SourceBlock] = []
    used_ocr = False
    with pymupdf.open(path) as document:
        metadata = document.metadata or {}
        title = str(metadata.get("title") or Path(original_filename).stem)
        for page_number, page in enumerate(document, start=1):
            page_blocks = page.get_text("blocks", sort=True)
            texts = [normalize_text(str(item[4])) for item in page_blocks if len(item) > 6 and int(item[6]) == 0]
            texts = [value for value in texts if value]
            page_used_ocr = False
            if not texts and enable_ocr:
                text_page = page.get_textpage_ocr(
                    language=ocr_language,
                    dpi=ocr_dpi,
                    full=True,
                )
                value = normalize_text(page.get_text("text", textpage=text_page, sort=True))
                texts = [value] if value else []
                page_used_ocr = bool(value)
                used_ocr = used_ocr or page_used_ocr
            for value in texts:
                blocks.append(
                    SourceBlock(
                        text=value,
                        block_type="ocr_text" if page_used_ocr else "paragraph",
                        page_start=page_number,
                        page_end=page_number,
                    )
                )
    return [
        _make_document(
            path,
            original_filename=original_filename,
            digest=digest,
            title=title,
            blocks=blocks,
            doc_type="pdf",
            metadata={"ocr_used": used_ocr},
        )
    ]


def _parse_excel(
    path: Path,
    original_filename: str,
    *,
    rows_per_block: int,
) -> list[KnowledgeDocument]:
    digest = _file_digest(path)
    suffix = Path(original_filename).suffix.lower()
    if suffix == ".xls":
        try:
            import pandas as pd
        except ImportError as exc:
            raise ValueError("XLS parsing requires pandas and xlrd") from exc
        sheets = pd.read_excel(path, sheet_name=None, dtype=object, header=None)
        iter_sheets = ((name, frame.fillna("").values.tolist()) for name, frame in sheets.items())
    else:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("Excel parsing requires openpyxl") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        iter_sheets = [
            (sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
            for sheet in workbook.worksheets
        ]
        workbook.close()

    documents: list[KnowledgeDocument] = []
    for sheet_index, (sheet_name, raw_rows) in enumerate(iter_sheets):
        rows = [row for row in raw_rows if any(value not in (None, "") for value in row)]
        if not rows:
            continue
        header = rows[0]
        blocks: list[SourceBlock] = []
        data_rows = rows[1:]
        if not data_rows:
            data_rows = [header]
        for start in range(0, len(data_rows), rows_per_block):
            part = data_rows[start : start + rows_per_block]
            table_rows = [header, *part] if part != [header] else [header]
            value = _table_markdown(table_rows)
            if value:
                blocks.append(
                    SourceBlock(
                        text=f"工作表：{sheet_name}\n{value}",
                        block_type="table",
                        section_path=[sheet_name],
                        row_start=start + 2,
                        row_end=start + len(part) + 1,
                        metadata={"sheet_name": sheet_name},
                    )
                )
        documents.append(
            _make_document(
                path,
                original_filename=original_filename,
                digest=digest,
                title=f"{Path(original_filename).stem} · {sheet_name}",
                blocks=blocks,
                doc_type="spreadsheet",
                id_suffix=f"sheet{sheet_index + 1}",
                metadata={"sheet_name": sheet_name},
            )
        )
    if not documents:
        raise ValueError(f"no readable worksheets found in {original_filename}")
    return documents


def parse_source(
    path: Path,
    *,
    original_filename: str | None = None,
    pdf_ocr_enabled: bool = False,
    pdf_ocr_language: str = "chi_sim+eng",
    pdf_ocr_dpi: int = 200,
    excel_rows_per_block: int = 25,
) -> list[KnowledgeDocument]:
    """按扩展名选择解析器并返回统一文档列表。"""
    filename = Path(original_filename or path.name).name
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"unsupported knowledge file type: {suffix or '<none>'}")
    if suffix == ".json":
        return _parse_json(path)
    if suffix == ".pdf":
        return _parse_pdf(
            path,
            filename,
            enable_ocr=pdf_ocr_enabled,
            ocr_language=pdf_ocr_language,
            ocr_dpi=pdf_ocr_dpi,
        )
    if suffix == ".docx":
        return _parse_docx(path, filename)
    if suffix in {".md", ".markdown"}:
        return _parse_markdown(path, filename)
    if suffix in {".html", ".htm"}:
        return _parse_html(path, filename)
    if suffix == ".txt":
        return _parse_txt(path, filename)
    return _parse_excel(path, filename, rows_per_block=excel_rows_per_block)
